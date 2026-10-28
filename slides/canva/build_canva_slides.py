# Builds slides/canva/family_table_3_slides.pptx: the three Canva pages owned by Rz
# (AI Client & Family Experience, AI Advisor Workbench, Added Value: The Secret Sauce & Core Innovation).
# Every box, label and table is a native, editable shape so it survives a Canva import; only app screens are images.
# Look matches the team deck (DAHY7Vfs02I): cream page, bold black titles, orange squares, slate table headers,
# an orange band at the bottom, a sources line. Screens come from video/src/slide_shots_canva.js.
# Run: python3 slides/canva/build_canva_slides.py
from pathlib import Path
from PIL import Image, ImageDraw, ImageFilter
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

HERE = Path(__file__).parent
RAW, ART = HERE / "raw", HERE / "art"
ART.mkdir(exist_ok=True)

CREAM, ORANGE, BAND, SLATE = "FFF3DE", "FF6A33", "FFB471", "344C5F"
INK, SUB, BLUE, WHITE, DARK = "111111", "4A4A4A", "0069AA", "FFFFFF", "05060D"
FONT = "Arial"
PX = 6350  # EMU per px on a 1920 x 1080 canvas


def rgb(h):
    return RGBColor.from_string(h)


# ---------- image framing ----------
def phone_frame(src, out, bezel=26, radius=96):
    """Screen inside a dark bezel with rounded corners; transparent outside."""
    im = Image.open(src).convert("RGBA")
    w, h = im.size
    W, H = w + 2 * bezel, h + 2 * bezel
    frame = Image.new("RGBA", (W, H), (0, 0, 0, 0))
    d = ImageDraw.Draw(frame)
    d.rounded_rectangle([0, 0, W - 1, H - 1], radius=radius + bezel, fill=(11, 21, 51, 255))
    mask = Image.new("L", (w, h), 0)
    ImageDraw.Draw(mask).rounded_rectangle([0, 0, w - 1, h - 1], radius=radius, fill=255)
    frame.paste(im, (bezel, bezel), mask)
    frame.save(out)


def window_frame(src, out, radius=22):
    """Desktop crop with rounded corners and a hairline border."""
    im = Image.open(src).convert("RGBA")
    w, h = im.size
    mask = Image.new("L", (w, h), 0)
    ImageDraw.Draw(mask).rounded_rectangle([0, 0, w - 1, h - 1], radius=radius, fill=255)
    out_im = Image.new("RGBA", (w, h), (0, 0, 0, 0))
    out_im.paste(im, (0, 0), mask)
    ImageDraw.Draw(out_im).rounded_rectangle([1, 1, w - 2, h - 2], radius=radius, outline=(214, 205, 190, 255), width=3)
    out_im.save(out)


for n in ["p1_preview", "p2_setup", "p3_checkin", "p4_decide", "p5_ask"]:
    phone_frame(RAW / f"{n}.png", ART / f"{n}.png")
for n in ["d1_signal", "d2_brief", "d3_guide", "d4_inbox"]:
    window_frame(RAW / f"{n}.png", ART / f"{n}.png")


# ---------- shape helpers (all positions in px) ----------
prs = Presentation()
prs.slide_width, prs.slide_height = Emu(1920 * PX), Emu(1080 * PX)
BLANK = prs.slide_layouts[6]


def box(s, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE, radius=None):
    sh = s.shapes.add_shape(shape, Emu(x * PX), Emu(y * PX), Emu(w * PX), Emu(h * PX))
    if fill:
        sh.fill.solid(); sh.fill.fore_color.rgb = rgb(fill)
    else:
        sh.fill.background()
    if line:
        sh.line.color.rgb = rgb(line); sh.line.width = Pt(1)
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = radius
    sh.text_frame.text = ""
    return sh


def text(s, x, y, w, h, runs, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.15, shape=None):
    """runs: a string, or a list of paragraphs; each paragraph a string or a list of (text, {overrides})."""
    tb = shape or s.shapes.add_textbox(Emu(x * PX), Emu(y * PX), Emu(w * PX), Emu(h * PX))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    paras = runs if isinstance(runs, list) else [runs]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        for t, o in ([(para, {})] if isinstance(para, str) else para):
            r = p.add_run(); r.text = t
            f = r.font; f.name = FONT; f.size = Pt(o.get("size", size) * 0.75)
            f.bold = o.get("bold", bold); f.color.rgb = rgb(o.get("color", color))
    return tb


def pic(s, path, x, y, w=None, h=None):
    return s.shapes.add_picture(str(path), Emu(x * PX), Emu(y * PX), Emu(w * PX) if w else None, Emu(h * PX) if h else None)


def chrome(s, title, sub, sources):
    """The team deck's frame: cream page, orange square cluster, bold title, subtitle with an orange rule, sources."""
    bg = s.background.fill; bg.solid(); bg.fore_color.rgb = rgb(CREAM)
    for x, y in [(1792, 0), (1728, 64), (1856, 64)]:
        box(s, x, y, 64, 64, ORANGE)
    text(s, 80, 54, 1560, 70, title, size=56, bold=True)
    text(s, 80, 136, 1560, 40, sub, size=24, color=SUB)
    box(s, 80, 186, 560, 4, ORANGE)
    text(s, 80, 1028, 1760, 30, sources, size=16, color=SUB)


def band(s, label, body, y=880, h=118):
    box(s, 0, y, 1920, h, BAND)
    text(s, 80, y, 320, h, label, size=24, bold=True, anchor=MSO_ANCHOR.MIDDLE, spacing=1.0)
    text(s, 400, y, 1440, h, body, size=21, anchor=MSO_ANCHOR.MIDDLE)


def pill(s, x, y, label, fill=ORANGE, color=WHITE, w=None, size=17):
    w = w or int(len(label) * size * 0.58 + 36)
    sh = box(s, x, y, w, 34, fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
    text(s, 0, 0, 0, 0, label, size=size, bold=True, color=color, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, shape=sh)
    return w


# ---------- slide 1: AI Client & Family Experience ----------
s = prs.slides.add_slide(BLANK)
chrome(s, "AI Client & Family Experience",
       "Two quick steps to start, then the core journey where the family decides: Align, Allocate, Activate",
       "Source: Family Table prototype (mobile). The Wilson family, portraits and figures are fictional; AI steps are simulated, formulas are real.")
cols = [
    ("Explore", "See what a sale could mean", "Before talking to anyone: the likely amount at closing, and two ways to split it.", "none; fixed ratios do the math.", "p1_preview"),
    ("Onboard", "Confirm facts, sharpen goals", "PNC fills in what it knows. Margaret ranks three goals and says what fair means.", "prefills from PNC records.", "p2_setup"),
    ("Align", "Speak in private", "Claire writes what she's never said, then chooses what Alex may read.", "NLP tags themes in her own vault.", "p3_checkin"),
    ("Allocate", "Vote, then 12 plans to 1", "Each child's number under every split, a private vote, and a decision tree.", "explains; formulas do every number.", "p4_decide"),
    ("Activate", "Ask a person, not a bot", "Questions go to Alex. The reply carries its numbers; a call is one tap away.", "drafts; Alex checks before it's sent.", "p5_ask"),
]
CW, GAP, X0 = 336, 20, 80
# The core journey (Align, Allocate, Activate) sits on one tinted panel; Explore and Onboard are the lead-in.
cx = X0 + 2 * (CW + GAP) - 16
box(s, cx, 200, 3 * CW + 2 * GAP + 32, 676, "FFE4CF", shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.03)
text(s, cx + 16, 208, 3 * CW + 2 * GAP, 26, "Core journey", size=17, bold=True, color=ORANGE, align=PP_ALIGN.RIGHT)
for i, (step, h3, body, ai, img) in enumerate(cols):
    x = X0 + i * (CW + GAP)
    if i < 2:
        pill(s, x, 214, step, fill="EADCC4", color=SLATE)
    else:
        pill(s, x, 214, step)
    pic(s, ART / f"{img}.png", x + (CW - 214) // 2, 262, w=214)  # phone ≈ 214 x 458
    text(s, x, 734, CW, 30, h3, size=22, bold=True)
    text(s, x, 766, CW, 62, body, size=17, color=SUB, spacing=1.1)
    text(s, x, 838, CW, 26, [[("AI: ", {"bold": True, "color": BLUE}), (ai, {"color": BLUE})]], size=16, spacing=1.05)
band(s, "WHAT THE FAMILY GETS",
     [[("Clarity: ", {"bold": True}), ("their own number, not a guess.   ", {}),
       ("Fairness: ", {"bold": True}), ("a private voice, without a fight at dinner.   ", {}),
       ("Confidence: ", {"bold": True}), ("every answer is checked by a person before it reaches them.", {})]])
s.notes_slide.notes_text_frame.text = (
    "Five moments on the family side, all on a phone. Explore and Onboard are the lead-in; the core journey, where the family actually "
    "decides, is Align, Allocate and Activate. Before Margaret ever talks to PNC, the preview shows what a sale "
    "could mean for a family like hers. Once she joins, PNC has already filled in the facts; she ranks her goals and says what fair "
    "means to her. Claire checks in privately: natural-language processing tags themes, and she decides what Alex may read. Each child "
    "sees their own number under every split, votes privately, and a decision tree narrows 12 possible plans to one. Questions go to "
    "Alex, not a chatbot: the AI drafts, Alex checks and sends. This answers Margaret's friction: a guided portal, not a dense dashboard, "
    "and no direct confrontation at the dinner table.")

# ---------- slide 2: AI Advisor Workbench ----------
s = prs.slides.add_slide(BLANK)
chrome(s, "AI Advisor Workbench",
       "Alex sees the sale a year early, then AI does the reading in each core stage and Alex makes every call",
       "Source: Family Table prototype. Scores, family data and figures are fictional; the signal model is a hypothesis to validate on PNC data.")
rows = [
    ("Explore", "Sale signal", "Machine learning scores business-account activity: 81 out of 100, 12–18 months before closing, with its reasons.", "Whether and when to call Margaret."),
    ("Align", "Family brief", "NLP turns four private check-ins into themes; a fixed formula scores the gaps (fairness 1.00); Document AI flags the 2012 will.", "Approves the brief. Books a Hawthorn facilitator when a gap passes 0.75."),
    ("Allocate", "Meeting guide and options", "Turns the brief into a 90-minute run of show and five presenting screens; explains each split, every number from the calculator.", "What to present. The facilitator leads the fairness talk; the family decides in principle."),
    ("Activate", "Inbox and specialists", "Drafts each reply from the plan; routes tasks to five specialist teams; flags $8.0–9.0M heading to heirs who bank elsewhere.", "Edits and sends every reply. Introduces PNC to Claire and Sam, with Margaret's agreement."),
]
TX, TW = 80, 1120
cw = [140, 200, 430, 350]
tbl = s.shapes.add_table(len(rows) + 1, 4, Emu(TX * PX), Emu(214 * PX), Emu(TW * PX), Emu(560 * PX)).table
for j, w in enumerate(cw):
    tbl.columns[j].width = Emu(w * PX)
heads = ["Stage", "AI tool", "What the AI does", "What Alex decides"]
for j, hd in enumerate(heads):
    c = tbl.cell(0, j); c.fill.solid(); c.fill.fore_color.rgb = rgb(SLATE)
    text(s, 0, 0, 0, 0, hd, size=19, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, shape=c)
tbl.rows[0].height = Emu(54 * PX)
for i, (stage, tool, ai, alex) in enumerate(rows, start=1):
    tbl.rows[i].height = Emu(124 * PX)
    core = stage != "Explore"
    for j, val in enumerate([stage, tool, ai, alex]):
        c = tbl.cell(i, j); c.fill.solid()
        c.fill.fore_color.rgb = rgb((ORANGE if core else "EADCC4") if j == 0 else ("FFFAF1" if i % 2 else "FFF3DE"))
        c.margin_left = c.margin_right = Emu(14 * PX); c.margin_top = c.margin_bottom = Emu(8 * PX)
        if j == 0:
            text(s, 0, 0, 0, 0, val, size=20, bold=True, color=(WHITE if core else SLATE), align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, shape=c)
        else:
            text(s, 0, 0, 0, 0, val, size=19 if j == 1 else 17, bold=(j == 1), color=SUB if j == 2 else INK, anchor=MSO_ANCHOR.MIDDLE, spacing=1.08, shape=c)
IX, IW = 1250, 590  # 1100x620 crops -> 590 x 333
pic(s, ART / "d2_brief.png", IX, 214, w=IW)
text(s, IX, 553, IW, 26, [[("AI brief: ", {"bold": True}), ("every AI step is shown, then Alex approves", {})]], size=16, color=SUB)
pic(s, ART / "d3_guide.png", IX, 590, w=IW // 2 - 6, h=int((IW // 2 - 6) * 620 / 1100))
pic(s, ART / "d4_inbox.png", IX + IW // 2 + 6, 590, w=IW // 2 - 6, h=int((IW // 2 - 6) * 620 / 1100))
text(s, IX, 760, IW, 50, [[("Meeting guide and inbox: ", {"bold": True}), ("AI drafts the run of show and each reply; Alex presents and sends", {})]], size=16, color=SUB, spacing=1.1)
band(s, "ADVISOR RULE",
     [[("AI reads, scores, drafts and routes. ", {}), ("Alex reviews every output before any client action,", {"bold": True}),
       (" and sees the family's gaps, never anyone's words.", {})]])
s.notes_slide.notes_text_frame.text = (
    "The advisor side, stage by stage: the signal in Explore, then one AI tool for each core stage, Align, Allocate and Activate. Each row says what the AI does and what Alex decides, because the rule is that Alex reviews every "
    "output before any client action. The signal comes from PNC's own business-account data, a year before closing. The brief is the "
    "core insight: Margaret told Alex the kids are on the same page, but the family's widest gap is fairness. Alex sees it as a range, "
    "with no names. The AI then turns the brief into a 90-minute meeting guide with screens to present, drafts replies to the family "
    "that Alex sends, and routes tasks to five specialist teams. It also flags the $8 to 9 million heading to two heirs who bank "
    "elsewhere. This answers Alex's friction: hearing only the founder, undocumented concerns, and capabilities spread across business units.")

# ---------- slide 3: Added Value — The Secret Sauce & Core Innovation ----------
s = prs.slides.add_slide(BLANK)
chrome(s, "Added Value: The Secret Sauce & Core Innovation",
       "Three assets only PNC can combine, and what they are worth to the family and to PNC",
       "Sources: Cerulli Associates (2019), 13% of adult children use their parents' advisor. Wilson figures are fictional; the fee is an illustrative 0.8% blended rate, to validate with PNC.")
pillars = [
    ("A signal PNC already holds", "Machine learning", "Powers Explore", "PNC banks the operating company. Payments to M&A advisers, audited-financials requests and early paydowns flag a sale 12–18 months ahead.",
     "Outside wealth managers see the money only after it moves.", DARK, WHITE, "C9D0DE"),
    ("Private-by-default family intelligence", "NLP + fixed formula", "Powers Align and Allocate", "Natural-language processing turns each member's private words into themes inside their own vault. A fixed formula scores where the family disagrees.",
     "Honest answers given in private can't be bought, scraped or rebuilt.", ORANGE, WHITE, "FFE6D9"),
    ("One household graph", "Consent ledger", "Powers Activate and Continue", "Business, owners, spouse, heirs, trusts and accounts in one consented map, under one roof with Business Banking, the Private Bank and Hawthorn IFS.",
     "Heirs become PNC relationships before any money moves.", "FFFFFF", INK, SUB),
]
PW, PG = 568, 28
for i, (h3, tag, powers, body, moat, fill, hc, bc) in enumerate(pillars):
    x = 80 + i * (PW + PG)
    box(s, x, 214, PW, 400, fill, line=("E6D9C3" if fill == "FFFFFF" else None), shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
    pill(s, x + 32, 244, tag, fill=(WHITE if fill != "FFFFFF" else BLUE), color=(fill if fill != "FFFFFF" else WHITE))
    text(s, x + 32, 296, PW - 64, 76, h3, size=30, bold=True, color=hc, spacing=1.05)
    text(s, x + 32, 380, PW - 64, 96, body, size=19, color=bc, spacing=1.15)
    text(s, x + 32, 484, PW - 64, 28, powers, size=18, bold=True, color=("F7A766" if fill == DARK else WHITE if fill == ORANGE else ORANGE))
    box(s, x + 32, 528, PW - 64, 2, ("3A4C82" if fill == DARK else "FFFFFF" if fill == ORANGE else "E6D9C3"))
    text(s, x + 32, 540, PW - 64, 60, [[("Hard to copy: ", {"bold": True}), (moat, {})]], size=18, color=hc, spacing=1.1)
# added value: the family and PNC
text(s, 80, 640, 560, 30, "For the family", size=22, bold=True, color=ORANGE)
fam = ["The fairness gap surfaces before the letter of intent, not after closing",
       "Every child sees their own number under every option",
       "A facilitator steps in when the gap is wide"]
text(s, 80, 676, 560, 190, [[("•  ", {"color": ORANGE, "bold": True}), (f, {})] for f in fam], size=19, spacing=1.25)
text(s, 700, 640, 1140, 30, "For PNC, in the Wilson case", size=22, bold=True, color=ORANGE)
nums = [("$27.0M", "proceeds stay at closing"), ("3 → 5", "family relationships at PNC"), ("$8.5M", "reaches the next generation at PNC"), ("~$216K", "a year in fees, at an illustrative 0.8%")]
NW = 1140 // 4
for i, (n, l) in enumerate(nums):
    x = 700 + i * NW
    text(s, x, 672, NW - 16, 70, n, size=54, bold=True, color=INK, spacing=1.0)
    text(s, x, 748, NW - 24, 60, l, size=18, color=SUB, spacing=1.1)
band(s, "CORE INNOVATION",
     [[("A consented family-intelligence layer on PNC's own commercial signal. ", {"bold": True}),
       ("AI reads and explains, formulas do the money, and the advisor decides, seeing how far apart the family is, never anyone's words, and only when 3 or more take part.", {})]])
s.notes_slide.notes_text_frame.text = (
    "Generative AI and alerts are easy to copy; these three assets are not. First, timing: PNC banks the operating company, so its own "
    "account data can flag a sale a year before closing, while outside wealth managers see the money only after it moves. Second, "
    "private-by-default family intelligence: natural-language processing turns each member's private words into themes inside their "
    "own vault, a fixed formula scores the gaps, and Alex sees the size of a disagreement, never the words, and only when at least "
    "three people take part. Honest answers given in private are data no competitor can buy or scrape. Third, one household graph "
    "links the business to every heir, so the next generation meets PNC before the money moves. The added value: for the family, the "
    "fairness question surfaces before the letter of intent. For PNC, in the Wilson case, $27.0M stays, relationships grow from three "
    "to five, $8.5M reaches the next generation at PNC, and the fees are about $216K a year at an illustrative 0.8%.")

out = HERE / "family_table_3_slides.pptx"
prs.save(out)
print("saved", out)
